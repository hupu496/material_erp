from django.shortcuts import render
from django.http import JsonResponse
from datetime import datetime
from masters.models import Category, SubCategory, WorkSpots, UnitMaster,PartyMaster
from materials.models import MaterialMaster,MaterialReceipt,MaterialRequisition,MaterialOutward
import base64,csv
from django.core.files.base import ContentFile
from openpyxl import Workbook

from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from django.http import HttpResponse
from io import BytesIO
from openpyxl.utils import get_column_letter
from django.db.models import Sum
from django.template.loader import get_template
from xhtml2pdf import pisa

from django.conf import settings
import os



def report_dashboard(request):

    if request.method == "POST":

        report_type = request.POST.get('report_type')
        category = request.POST.get('category')
        subcategory = request.POST.get('subcategory')
        if report_type == 'party':

            queryset = PartyMaster.objects.select_related(
                'category',
            ).all()
            if category:
                queryset = queryset.filter(category_id=category)

            filename = f"Party_Report_{datetime.now().strftime('%Y%m%d')}"
            return generate_report(
                queryset,
                'party',
                filename,
                request.POST.get('format', 'xlsx')
            )

        elif report_type == 'material':

            queryset = MaterialMaster.objects.select_related(
                'category',
                'sub_category'
            ).all()

            if category:
                queryset = queryset.filter(category_id=category)

            if subcategory:
                queryset = queryset.filter(sub_category_id=subcategory)

            filename = f"Material_Report_{datetime.now().strftime('%Y%m%d')}"

            return generate_report(
                queryset,
                'material',
                filename,
                request.POST.get('format', 'xlsx')
            )

    categories = Category.objects.all()
   
    return render(request, 'reports/master_report.html', {
        'categories': categories,
    })



def generate_report(queryset, report_for, filename, file_format='xlsx'):
    if file_format == 'pdf':

        context = {
            'report_title': filename.replace('_', ' '),
            'generated_on': datetime.now(),
            'generated_by': 'Admin',
            'report_for': report_for,
            'rows': queryset,
            'logo_path': os.path.join(
                settings.BASE_DIR,
                'static/images/images.png'
            )
        }

        template = get_template(
            'reports/master_report_pdf.html'
        )

        html = template.render(context)

        response = HttpResponse(
            content_type='application/pdf'
        )

        response[
            'Content-Disposition'
        ] = f'attachment; filename="{filename}.pdf"'

        pisa_status = pisa.CreatePDF(
            html,
            dest=response
        )

        if pisa_status.err:
            return HttpResponse(
                'PDF generation failed'
            )

        return response

    if file_format == 'csv':

        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = \
            f'attachment; filename="{filename}.csv"'

        writer = csv.writer(response)

        # PARTY REPORT
        if report_for == 'party':

            writer.writerow([
                'Party Name',
                'Address',
                'Phone',
                'Email',
                'Category',
                
            ])

            for item in queryset:

                writer.writerow([
                    item.name,
                    item.address,
                    item.phone,
                    item.email,
                    item.category.name if item.category else '',
                   
                ])

        # MATERIAL REPORT
        elif report_for == 'material':

            writer.writerow([
                'ID',
                'Description',
                'Category',
                'Sub Category',
                'Work Spot',
                'Units',
                'Rate',
                'Quantity',
                'Rack No',
                'Total Stock',
                'Critical Qty',
                
            ])

            for item in queryset:

                writer.writerow([
                    item.id,
                    item.description,
                    item.category.name if item.category else '',
                    item.sub_category.sub_category if item.sub_category else '',
                    item.work_spots.name if item.work_spots else '',
                    item.units.name if item.units else '',
                    item.rate,
                    item.quantity,
                    item.rackno,
                    item.qty,
                    item.critical_qty
                ])

        return response

    else:

        wb = Workbook()
        ws = wb.active
        ws.title = "Master Report"
        if report_for == 'party':
            ws.title = "INDIAN OIL CORPORATION LIMITED - PARTY MASTER REPORT"
            print(ws.title)
        else:
            ws.title = "INDIAN OIL CORPORATION LIMITED - MATERIAL MASTER REPORT"
        

        # Merge title cells
        ws.merge_cells('A1:K1')

        # Set title
        ws['A1'] = ws.title

        # Style title
        ws['A1'].font = Font(bold=True, size=16)
        ws['A1'].alignment = Alignment(horizontal='center')

        # Report generated date
        ws.merge_cells('A2:K2')
        ws['A2'] = f"Generated On: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}"
        ws['A2'].alignment = Alignment(horizontal='center')

        # Empty row
        ws.append([])

        # PARTY REPORT
        if report_for == 'party':

            headers = [
                'Party Name',
                'Address',
                'Phone',
                'Email',
                'Category',
               
            ]
            ws.append(headers)
            for item in queryset:

                ws.append([
                    item.name,
                    item.address,
                    item.phone,
                    item.email,
                    item.category.name if item.category else '',
                  
                ])

        # MATERIAL REPORT
        elif report_for == 'material':

            headers = [
                'ID',
                'Description',
                'Category',
                'Sub Category',
                'Work Spot',
                'Units',
                'Rate',
                'Quantity',
                'Rack No',
                'Total Stock',
                'Critical Qty',
                
            ]

            ws.append(headers)

            for item in queryset:
                ws.append([
                    item.id,
                    item.description,
                    item.category.name if item.category else '',
                    item.sub_category.sub_category if item.sub_category else '',
                    item.work_spots.name if item.work_spots else '',
                    item.units.name if item.units else '',
                    item.rate,
                    item.quantity,
                    item.rackno,
                    item.qty,
                    item.critical_qty
                ])
        for col in ws.iter_cols(min_row=1, max_row=1):
            for cell in col:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        response['Content-Disposition'] = \
            f'attachment; filename="{filename}.xlsx"'
        wb.save(response)
        return response
def load_subcategories(request):

    category_id = request.GET.get('category')

    subcategories = SubCategory.objects.filter(
        category_id=category_id
    ).values('id', 'sub_category')

    return JsonResponse(list(subcategories), safe=False)



def date_wise_ledger(request):

    categories = Category.objects.filter(status=True)

    queryset = None
    report_data = []

    if request.method == "POST":

        ledger_type = request.POST.get('ledger_type')
        from_date = request.POST.get('from_date')
        to_date = request.POST.get('to_date')
        category = request.POST.get('category')
        file_type = request.POST.get('file_type')


        # ================= RECEIPT =================

        if ledger_type == "receipt":

            queryset = MaterialReceipt.objects.all()
            if from_date and to_date:
                queryset = queryset.filter(
                    issue_date__range=[from_date, to_date]
                )

        # ================= REQUISITION =================

        elif ledger_type == "requisition":

            queryset = MaterialRequisition.objects.all()

            if from_date and to_date:

                queryset = queryset.filter(
                    issue_date__range=[from_date, to_date]
                )

        # ================= CATEGORY FILTER =================

        if category:

            queryset = queryset.filter(
                category_id=category
            )

        # ================= CSV EXPORT =================

        if file_type == "csv":

            response = HttpResponse(
                content_type='text/csv'
            )

            response['Content-Disposition'] = \
                'attachment; filename="ledger.csv"'

            writer = csv.writer(response)

            writer.writerow([
                'Date',
                'Entry No',
                'Category',
                'Item',
                'Qty'
            ])

            for item in queryset:

                qty = item.total_stock \
                    if ledger_type == "receipt" \
                    else item.total_nos

                writer.writerow([
                    item.issue_date,
                    item.entry_no,
                    item.category,
                    item.item,
                    qty
                ])

            return response

        # ================= EXCEL EXPORT =================

        elif file_type == "xlsx":

            wb = Workbook()
            ws = wb.active

            ws.title = "Ledger Report"

            ws.append([
                'Date',
                'Entry No',
                'Category',
                'Item',
                'Qty'
            ])

            for item in queryset:

                qty = item.total_stock \
                    if ledger_type == "receipt" \
                    else item.total_nos

                ws.append([
                    str(item.issue_date),
                    str(item.entry_no),
                    str(item.category),
                    str(item.item),
                    str(qty)
                ])

            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )

            response['Content-Disposition'] = \
                'attachment; filename="ledger.xlsx"'

            wb.save(response)

            return response

        # ================= PDF PREVIEW =================

        elif file_type == "pdf":

            total_qty = 0
            sl = 1

            for item in queryset:

                if ledger_type == "receipt":

                    qty = MaterialMaster.objects.get(id=item.item_id)
                    unit= UnitMaster.objects.get(id=qty.units_id)
                  
                    vendor = PartyMaster.objects.get(id=item.receive_from)
                    total_qtys=item.totalno
                    report_data.append({

                        'sl_no': sl,
                        'date': item.issue_date,
                        'entry_no': item.entry_no,
                        'challan_no': item.challan_no,
                        'category': item.category,
                        'item': item.item.description,
                        'opening_qty': item.total_stock,
                        'in_use_qty': qty.qty if qty else '',
                        'unit': unit.name if unit else '',
                        'vendor': vendor.name if vendor else '',
                        'vehicle': item.vehicle_no,
                        'receiver': item.receiver,
                        'received_by': item.receiver,
                        'receipt_qty': total_qtys,
                        'remarks': item.remarks,
                    })

                else:
                    qty = MaterialMaster.objects.get(id=item.item_id)
                    unit= UnitMaster.objects.get(id=qty.units_id)
                    workspot = WorkSpots.objects.get(id=item.work_spots_id)
                    total_qtys=item.requisition_nos

                    report_data.append({

                        'sl_no': sl,
                        'date': item.issue_date,
                        'entry_no': item.entry_no,
                        'challan_no': item.requisition_no,
                        'category': item.category,
                        'item': item.item.description,
                        'opening_qty': item.total_nos,
                        'in_use_qty': qty.qty if qty else '',
                        'unit': unit.name if unit else '',
                        'WorkSpot': workspot.name if workspot else '',
                        'vehicle': '-',
                        'receiver': item.requisition_by,
                        'received_by': item.requisition_by,
                        'receipt_qty': item.requisition_nos,
                        'remarks': item.remarks,
                    })

                total_qty += int(total_qtys)
                sl += 1

            context = {

                'report_data': report_data,
                'total_receipt_qty': total_qty,
                'company_name': 'Indian Oil Corporation Ltd.',

                'plant_name': 'Material Inventory System',

                'report_title': 'Date Wise Ledger Report',
                'type' : ledger_type,

                'from_date': from_date,

                'to_date': to_date,

                'generated_on': datetime.now(),

                'generated_by': request.user.username,
            }

            return render(
                request,
                'reports/material_receipt_ledger_pdf.html',
                context
            )

    return render(
        request,
        'reports/date_wise_ledger.html',
        {
            'categories': categories
        }
    )

def stock_summary_form(request):
    categories = Category.objects.all()
    subcategories = SubCategory.objects.all()

    return render(request, "reports/stock_summary_form.html", {
        "categories": categories,
        "subcategories": subcategories,
    })
def stock_summary_report(request):

    categories = Category.objects.all()
    subcategories = SubCategory.objects.all()

    data = []
    from_date = request.GET.get("from_date")
    to_date = request.GET.get("to_date")
    category_id = request.GET.get("category")
    sub_category_id = request.GET.get("sub_category")

    if from_date and to_date:

        from_date = datetime.strptime(from_date, "%Y-%m-%d").date()
        to_date = datetime.strptime(to_date, "%Y-%m-%d").date()

        items = MaterialMaster.objects.all()

        if category_id:
            items = items.filter(category_id=category_id)

        if sub_category_id:
            items = items.filter(sub_category_id=sub_category_id)
        
        for item in items:
            # ---------------- OPENING RECEIPT ----------------
            opening_receipt = MaterialReceipt.objects.filter(
                item=item,
                issue_date__lt=from_date
            ).aggregate(total=Sum('totalno'))['total'] or 0

            # ---------------- OPENING ISSUE ----------------
            opening_issue = MaterialRequisition.objects.filter(
                item=item,
                issue_date__lt=from_date
            ).aggregate(total=Sum('requisition_nos'))['total'] or 0

            opening_qty = item.quantity

            # ---------------- PERIOD RECEIPT ----------------
            period_receipt = MaterialReceipt.objects.filter(
                item=item,
                issue_date__range=(from_date, to_date)
            ).aggregate(total=Sum('totalno'))['total'] or 0

            # ---------------- PERIOD ISSUE ----------------
            period_issue = MaterialRequisition.objects.filter(
                item=item,
                issue_date__range=(from_date, to_date)
            ).aggregate(total=Sum('requisition_nos'))['total'] or 0

            # ---------------- CLOSING ----------------
            closing_qty = opening_qty + period_receipt - period_issue
            
            data.append({
                "category": item.category.name,
                "sub_category": item.sub_category.sub_category,
                "description": item.description,
                "unit": item.units.name,
                "critical_qty": item.critical_qty,
                "opening_qty": item.quantity,
                "receipt": period_receipt,
                "issue": period_issue,
                "closing_qty": closing_qty,
            })

    return render(request, "reports/stock_summary.html", {
        "data": data,
        "categories": categories,
        "subcategories": subcategories,
    })